from kivy.uix.recyclegridlayout import defaultdict
from kivymd.app import MDApp
from kivy.uix.screenmanager import SlideTransition, NoTransition
from kivy.properties import StringProperty, NumericProperty
from kivymd.uix.screen import MDScreen
from kivymd.uix.boxlayout import MDBoxLayout
from kivy.core.window import Window
from kivy.uix.image import Image as KivyImage
from kivy.animation import Animation
from kivy.properties import ListProperty, DictProperty, ObjectProperty, BooleanProperty
from kivy.metrics import dp

from kivymd.uix.card import MDCard
from kivymd.uix.menu import MDDropdownMenu
from kivymd.uix.label import MDLabel
from kivymd.uix.textfield import MDTextField
from kivymd.uix.dialog import MDDialog
from kivymd.uix.filemanager import MDFileManager
from kivymd.uix.button import MDFlatButton, MDRaisedButton, MDIconButton
import os
import db_manager
from kivymd.uix.snackbar import Snackbar

# Window.size = (350, 600)

class BaseScreen(MDScreen):
    pass

class CartListItem(MDBoxLayout):
    item_name = StringProperty()
    item_price = NumericProperty()
    item_quantity = NumericProperty()
    actual_item_image_path = StringProperty("atlas://data/images/defaulttheme/transparent") 

    def __init__(self, item_data, **kwargs):
        super().__init__(**kwargs)
        self.item_name = item_data['name']
        self.item_price = item_data['price']
        self.item_quantity = item_data['quantity']

        _image_path = item_data.get('image_path', "images/placeholder.png")
        placeholder_path = "images/placeholder.png" 

        if _image_path and os.path.exists(_image_path):
            self.actual_item_image_path = _image_path
        elif os.path.exists(placeholder_path):
            self.actual_item_image_path = placeholder_path
        else: 
            self.actual_item_image_path = "atlas://data/images/defaulttheme/image-missing" 


    def increment_quantity(self, *args):
        app = MDApp.get_running_app()
        if app:
            app.update_cart_item_quantity(self.item_name, 1)
        else:
            print(f"[ERROR]: Aplikasi tidak ditemukan (CartListItem).")

    def decrement_quantity(self, *args):
        app = MDApp.get_running_app()
        if app:
            app.update_cart_item_quantity(self.item_name, -1)
        else:
            print(f"[ERROR]: Aplikasi tidak ditemukan (CartListItem).")

class Cart(MDScreen):
    payment_panel_visible = BooleanProperty(False)
    
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
    
    def on_enter(self):
        self.update_cart_display()

    def update_cart_display(self):
        app = MDApp.get_running_app()
        if not app or not hasattr(self, 'ids') or 'cart_items_list' not in self.ids:
            return

        cart_list_widget = self.ids.cart_items_list
        cart_list_widget.clear_widgets()
        cart_items_data = app.get_cart_items_list()

        if not cart_items_data:
            cart_list_widget.add_widget(
                MDLabel(
                    text="Keranjang Anda kosong.", halign="center",
                    theme_text_color="Secondary", padding_y=dp(20),
                    font_style="Subtitle1", adaptive_height=True
                )
            )
            if self.payment_panel_visible: 
                self.toggle_payment_panel() 
        else:
            for item_data_dict in cart_items_data:
                cart_item_view = CartListItem(item_data=item_data_dict)
                cart_list_widget.add_widget(cart_item_view)

        total_price = app.get_cart_total()
        total_items_count = sum(item_d.get('quantity', 0) for item_d in cart_items_data)

        if hasattr(self, 'ids'):
            if 'total_price_label' in self.ids:
                self.ids.total_price_label.text = f"Rp{total_price:,.0f}".replace(",", ".")
            if 'total_items_label' in self.ids:
                self.ids.total_items_label.text = str(total_items_count)
            
            payment_field = self.ids.get('payment_amount_field') 
            if payment_field:
                payment_field.text = ""  
            change_label = self.ids.get('change_label') 
            if change_label: 
                change_label.text = "Rp0"

            pay_trigger_button = self.ids.get('pay_trigger_button')
            if pay_trigger_button:
                pay_trigger_button.disabled = not cart_items_data
                if not cart_items_data:
                    pay_trigger_button.text = "KERANJANG KOSONG"
                    pay_trigger_button.icon = "cart-off"
                elif self.payment_panel_visible: 
                    pay_trigger_button.text = "TUTUP PEMBAYARAN"
                    pay_trigger_button.icon = "chevron-double-down"
                else: 
                    pay_trigger_button.text = "LANJUT KE PEMBAYARAN"
                    pay_trigger_button.icon = "chevron-double-up"


    def toggle_payment_panel(self):
        panel = self.ids.get('payment_details_panel')
        trigger_button = self.ids.get('pay_trigger_button')
        scrollview = self.parent.ids.get('cart_items_scrollview') if self.parent and hasattr(self.parent, 'ids') else None

        if not panel or not trigger_button:
            return

        duration = 0.3 

        if self.payment_panel_visible:
            anim = Animation(height=0, opacity=0, d=duration, t='out_cubic')
            trigger_button.text = "LANJUT KE PEMBAYARAN"
            trigger_button.icon = "chevron-double-up"
        else:
            target_height = dp(200)
            anim = Animation(height=target_height, opacity=1, d=duration, t='in_out_cubic')
            trigger_button.text = "TUTUP PEMBAYARAN"
            trigger_button.icon = "chevron-double-down"
            # Scroll ke bawah agar panel pembayaran terlihat
            if scrollview:
                def scroll_to_panel(*_):
                    scrollview.scroll_y = 0
                anim.bind(on_complete=lambda *_: scroll_to_panel())
        anim.start(panel)
        self.payment_panel_visible = not self.payment_panel_visible

    def calculate_and_display_change(self):
        app = MDApp.get_running_app()
        if not app or not hasattr(self, 'ids'):
            return

        total_price = app.get_cart_total()
        payment_amount_field = self.ids.get('payment_amount_field')
        change_label = self.ids.get('change_label')

        if not (payment_amount_field and change_label):
            return

        payment_amount_text = payment_amount_field.text
        if not payment_amount_text:
            change_label.text = "Rp0"
            return

        try:
            payment_amount = float(payment_amount_text)
            if hasattr(payment_amount_field, 'error'): payment_amount_field.error = False
        except ValueError:
            change_label.text = "Input tidak valid"
            if hasattr(payment_amount_field, 'error'): payment_amount_field.error = True
            return

        if payment_amount >= total_price:
            change = payment_amount - total_price
            change_label.text = f"Rp{change:,.0f}".replace(",", ".")
        else:
            change_label.text = "Rp0"

    def process_payment(self):
        app = MDApp.get_running_app()
        if not app or not hasattr(self, 'ids'):
            return

        total_price = app.get_cart_total()
        payment_amount_field = self.ids.get('payment_amount_field')
        
        if not payment_amount_field:
            return

        payment_amount_text = payment_amount_field.text
        if not payment_amount_text:
            if hasattr(payment_amount_field, 'error'): payment_amount_field.error = True
            if hasattr(payment_amount_field, 'helper_text'): payment_amount_field.helper_text = "Nominal tidak boleh kosong"
            return
        
        try:
            payment_amount = float(payment_amount_text)
            if hasattr(payment_amount_field, 'error'): payment_amount_field.error = False
            if hasattr(payment_amount_field, 'helper_text'): payment_amount_field.helper_text = "" 
        except ValueError:
            if hasattr(payment_amount_field, 'error'): payment_amount_field.error = True
            if hasattr(payment_amount_field, 'helper_text'): payment_amount_field.helper_text = "Input harus angka"
            return

        if payment_amount < total_price:
            if hasattr(payment_amount_field, 'error'): payment_amount_field.error = True
            if hasattr(payment_amount_field, 'helper_text'): payment_amount_field.helper_text = "Jumlah pembayaran kurang dari total"
            return
        
        change = payment_amount - total_price
        print(f"Pembayaran Berhasil! Total: {total_price}, Dibayar: {payment_amount}, Kembalian: {change}")

        
        outlet_name = app.get_active_outlet() or "Unknown"
        items = app.get_cart_items_list()
        db_manager.save_sale(outlet_name, items, total_price)
        
        if hasattr(app, 'cart_items'): 
            app.cart_items.clear()

        if self.payment_panel_visible:
            self.toggle_payment_panel()
        self.update_cart_display() 

class MenuCard(MDCard):
    name = StringProperty()
    image_path = StringProperty()
    price = NumericProperty()
    actual_image_source = StringProperty()

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.bind(image_path=self._update_actual_image_source)
        self._update_actual_image_source()

    def _update_actual_image_source(self, *args):
        placeholder = "images/placeholder.png"
        if self.image_path and os.path.exists(self.image_path):
            self.actual_image_source = self.image_path
        elif os.path.exists(placeholder):
            self.actual_image_source = placeholder 
        else:
            self.actual_image_source = ""

    def add_item_to_cart(self):
        app = MDApp.get_running_app()
        if not app:
            return

        app.add_item_to_app_cart(self)
        
        try:
            if not self.actual_image_source:
                return

            source_visual_element = self.ids.product_image_in_card
            if not source_visual_element:
                return
            
            start_pos_abs = source_visual_element.to_window(*source_visual_element.center)

            cart_icon_widget = app.root.ids.get('cart_nav_icon_target')
            if not cart_icon_widget: 
                return

            target_pos_abs = cart_icon_widget.to_window(*cart_icon_widget.center)

            anim_item = KivyImage(
                source=self.actual_image_source,
                size_hint=(None, None),
                size=(dp(40), dp(40)),
                allow_stretch=True,
                keep_ratio=True,
                mipmap=True
            )
            anim_item.center_x = start_pos_abs[0]
            anim_item.center_y = start_pos_abs[1]

            Window.add_widget(anim_item) 

            duration = 0.5
            anim = Animation(
                center_x=target_pos_abs[0],
                center_y=target_pos_abs[1],
                size=(dp(10), dp(10)),
                opacity=0.3,
                t='out_quad',
                duration=duration
            )
            
            def on_animation_finish(animation, widget_animated):
                Window.remove_widget(widget_animated)
            
            anim.bind(on_complete=on_animation_finish)
            anim.start(anim_item)

        except Exception as e:
            print(f"Animasi gagal: {e}")


class AddItemDialogContent(MDBoxLayout):
    file_manager = None
    manager_open = False

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.orientation = "vertical"
        self.spacing = dp(12)
        self.size_hint_y = None
        self.adaptive_height = True
        self.padding = [dp(16), dp(50), dp(16), dp(24)]

        self.name_field = MDTextField(hint_text="Nama Menu", mode="fill", required=True) 

        image_input_layout = MDBoxLayout(orientation='horizontal', adaptive_height=True, spacing=dp(10)) 
        self.image_path_display = MDTextField(
            hint_text="Pilih Gambar",
            mode="fill",
            readonly=True,
            helper_text_mode = "on_focus"
        )
        self.browse_button = MDIconButton(
            icon="folder-open-outline", 
            on_release=self.open_file_manager,
            pos_hint={'center_y': 0.5} 
        )
        image_input_layout.add_widget(self.image_path_display)
        image_input_layout.add_widget(self.browse_button)

        self.price_field = MDTextField(hint_text="Harga", input_filter="float", mode="fill", required=True)

        self.add_widget(self.name_field)
        self.add_widget(image_input_layout)
        self.add_widget(self.price_field)

    def open_file_manager(self, *args):
        from kivy.utils import platform
        try:
            if platform == 'android':
                from android.permissions import Permission, check_permission, request_permissions # type: ignore
                def after_permission(*_):
                    from android.storage import primary_external_storage_path # type: ignore
                    path = primary_external_storage_path()
                    if not os.path.exists(path):
                        self.show_error_dialog("Folder tidak ditemukan atau tidak dapat diakses.")
                        return
                    self._show_file_manager(path)
                # Cek apakah sudah diizinkan
                if check_permission(Permission.READ_EXTERNAL_STORAGE) and check_permission(Permission.WRITE_EXTERNAL_STORAGE):
                    after_permission()
                else:
                    def cb(permissions, results):
                        if all(results):
                            after_permission()
                        else:
                            self.show_error_dialog("Izin akses storage ditolak.")
                    request_permissions(
                        [Permission.READ_EXTERNAL_STORAGE, Permission.WRITE_EXTERNAL_STORAGE],
                        cb
                    )
                return
            else:
                path = os.path.expanduser("~")
                if not os.path.exists(path):
                    self.show_error_dialog("Folder home tidak ditemukan.")
                    return
            self._show_file_manager(path)
        except Exception as e:
            self.show_error_dialog(f"Gagal membuka file manager:\n{e}")
            
    def show_error_dialog(self, message):
        dialog = MDDialog(
            title="Gagal Membuka File Manager",
            text=message,
            size_hint=(0.8, None),
            height=dp(150),
            buttons=[
                MDFlatButton(
                    text="TUTUP",
                    on_release=lambda x: dialog.dismiss()
                )
            ]
        )
        dialog.open()

    def _show_file_manager(self, path):
        AddItemDialogContent.file_manager = MDFileManager(
            exit_manager=self.exit_manager,
            select_path=self.select_path,
            preview=True,
            ext=['.png', '.jpg', '.jpeg', '.gif', '.bmp']
        )
        AddItemDialogContent.file_manager.show(path)
        AddItemDialogContent.manager_open = True

    def select_path(self, path: str):
        self.image_path_display.text = path
        self.exit_manager()

    def exit_manager(self, *args):
        if AddItemDialogContent.file_manager:
            AddItemDialogContent.file_manager.close()
        AddItemDialogContent.manager_open = False

class AddMenu(MDScreen):
    dialog = None

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.all_menu_items_from_db = []
    
    def on_enter(self):
        self.load_items_from_db()
        search_input_widget = self.ids.get('search_field')
        if search_input_widget and search_input_widget.text:
            self.search_items(search_input_widget.text)
        else:
            self.display_all_items()
    
    def load_items_from_db(self):
        self.all_menu_items_from_db = db_manager.get_all_menu_items()

    def display_all_items(self):
        if not hasattr(self, 'ids') or not self.ids:
            return
        result_list = self.ids.get('search_result_list')
        if not result_list:
            return
        
        result_list.clear_widgets()
        for item_data in self.all_menu_items_from_db:
            self._create_and_add_card(item_data, result_list)

    def search_items(self, query):
        if not hasattr(self, 'ids') or not self.ids:
            return
        
        result_list = self.ids.get('search_result_list')
        if not result_list:
            return
        result_list.clear_widgets()
        if not query:
            self.display_all_items()
            return 

        for item_data in self.all_menu_items_from_db:
            if query.lower() in item_data['name'].lower():
                self._create_and_add_card(item_data, result_list)

    def _create_and_add_card(self, item_data, result_list_widget):
        card = MenuCard(
            name=item_data['name'],
            image_path=item_data.get('image', ''),
            price=item_data['price']
        )
        result_list_widget.add_widget(card)

    def open_add_item_dialog(self):
        if not self.dialog:
            content = AddItemDialogContent()
            self.dialog = MDDialog(
                title="Tambah Menu Baru",
                type="custom",
                content_cls=content,
                buttons=[
                    MDFlatButton(
                        text="BATAL",
                        on_release=self.dismiss_dialog
                    ),
                    MDRaisedButton(
                        text="TAMBAH",
                        on_release=self.submit_new_item_from_dialog
                    ),
                ],
            )
        dialog_actual_content = self.dialog.content_cls 

        dialog_actual_content.name_field.text = ""
        dialog_actual_content.image_path_display.text = ""
        dialog_actual_content.price_field.text = ""
        dialog_actual_content.name_field.error = False
        dialog_actual_content.price_field.error = False
        self.dialog.open()

    def dismiss_dialog(self, *args):
        if self.dialog:
            self.dialog.dismiss()

    def submit_new_item_from_dialog(self, *args):
        if not self.dialog:
            print("Error: Dialog tidak ditemukan di submit_new_item_from_dialog")
            return
        
        content = self.dialog.content_cls

        name = content.name_field.text
        image_path = content.image_path_display.text
        price_text = content.price_field.text
        
        content.name_field.error = False
        content.price_field.error = False

        if not name or not price_text:
            if not name: content.name_field.error = True
            if not price_text: content.price_field.error = True
            return
        
        try:
            price = float(price_text)
            db_manager.add_menu_item(name, image_path, price)
            self.load_items_from_db()

            current_search_text = self.ids.get('search_field').text if self.ids.get('search_field') else ""
            if current_search_text:
                self.search_items(current_search_text)
            else:
                self.display_all_items()
            
            if self.dialog:
                self.dialog.dismiss()
        except ValueError:
            content.price_field.error = True
        except Exception as e:
            print(f"Gagal menambahkan item: {e}")

class AddTableDialogContent(MDBoxLayout):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.orientation = "vertical"
        self.spacing = dp(12)
        self.size_hint_y = None
        self.adaptive_height = True
        self.padding = [dp(24), dp(24), dp(24), dp(24)]
        
        self.table_name_field = MDTextField(
            hint_text="No Meja",
            mode="fill",
            required=True,
            helper_text_mode="on_error"
        )
        self.add_widget(self.table_name_field)

class TableHistoryListItem(MDBoxLayout):
    table_data = DictProperty()  
    history_screen = ObjectProperty()  

    item_name = StringProperty("")
    item_status = StringProperty("")

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.bind(table_data=self._update_text_from_data)
        if self.table_data:
            self._update_text_from_data()

    def _update_text_from_data(self, *args):
        self.item_name = self.table_data.get('name', '')
        self.item_status = f"Status: {self.table_data.get('status', '')}"

    def show_options_menu(self, button_instance):
        if self.history_screen:
            self.history_screen.open_table_options_menu(self.table_data, button_instance)


class History(MDScreen):
    add_table_dialog = None
    table_options_dropdown = None 
    tables_data = ListProperty([]) 

    def on_enter(self):
        self.tables_data = db_manager.get_all_tables()
        self.display_tables()

    def display_tables(self):
        history_list_widget = self.ids.get('history_list')
        if not history_list_widget:
            return
        
        history_list_widget.clear_widgets() 

        if not self.tables_data:
            history_list_widget.add_widget(
                MDLabel(
                    text="Belum ada histori meja.", halign="center",
                    theme_text_color="Secondary", padding_y=dp(20),
                    font_style="Subtitle1", adaptive_height=True
                )
            )
        else:
            for table_info_original in sorted(self.tables_data, key=lambda x: x['name']):
                table_info = dict(table_info_original)  
                
                item = TableHistoryListItem(
                    table_data=table_info,
                    history_screen=self
                )
                history_list_widget.add_widget(item)
    
    def open_table_options_menu(self, table_info_dict, caller_button):
        table_name = table_info_dict['name']
        
        menu_items = [
            {
                "text": "Sudah Bayar",
                "viewclass": "OneLineListItem",
                "on_release": lambda x="Sudah Bayar": self.set_table_status(table_name, x),
            },
            {
                "text": "Belum Bayar",
                "viewclass": "OneLineListItem",
                "on_release": lambda x="Belum Bayar": self.set_table_status(table_name, x),
            },
            {
                "text": "Kosong",
                "viewclass": "OneLineListItem",
                "on_release": lambda x="Kosong": self.set_table_status(table_name, x),
            }
        ]
        
        if self.table_options_dropdown and self.table_options_dropdown.menu.parent:
            self.table_options_dropdown.dismiss()

        self.table_options_dropdown = MDDropdownMenu(
            caller=caller_button,
            items=menu_items,
            width_mult=2.5,
            position="center",
            max_height=dp(180),
        )
        self.table_options_dropdown.open()

    def set_table_status(self, table_name, new_status):
        success = db_manager.update_table_status(table_name, new_status)
        if success:
            self.tables_data = db_manager.get_all_tables()  
            self.display_tables()  
        
        if self.table_options_dropdown: 
            self.table_options_dropdown.dismiss()

    def open_add_table_dialog(self):
        if not self.add_table_dialog:
            content = AddTableDialogContent()
            self.add_table_dialog = MDDialog(
                title="Tambah Meja Baru",
                type="custom",
                content_cls=content, 
                buttons=[
                    MDFlatButton(text="BATAL", on_release=lambda x: self.add_table_dialog.dismiss()), #type: ignore
                    MDRaisedButton(text="TAMBAH", on_release=self.submit_new_table_from_dialog),
                ],
            )
        
        self.add_table_dialog.content_cls.table_name_field.text = ""
        self.add_table_dialog.content_cls.table_name_field.error = False
        self.add_table_dialog.content_cls.table_name_field.helper_text = ""
        self.add_table_dialog.open()

    def submit_new_table_from_dialog(self, *args):
        content_cls = self.add_table_dialog.content_cls #type: ignore
        table_name = content_cls.table_name_field.text.strip() 

        if not table_name:
            content_cls.table_name_field.error = True
            content_cls.table_name_field.helper_text = "Nama meja tidak boleh kosong."
            return

        existing_tables = db_manager.get_all_tables()
        if any(t['name'].lower() == table_name.lower() for t in existing_tables):
            content_cls.table_name_field.error = True
            content_cls.table_name_field.helper_text = "Nama meja sudah ada."
            return
        
        new_table_id = db_manager.add_table(table_name, 'Kosong')
        
        if new_table_id:
            self.tables_data = db_manager.get_all_tables()  
            self.display_tables()  
            self.add_table_dialog.dismiss() #type: ignore
        else:
            content_cls.table_name_field.error = True
            content_cls.table_name_field.helper_text = "Gagal menambahkan meja ke database."

    
    def open_status_update_dialog(self, table_info_dict):
        self.selected_table_for_status_update = table_info_dict  

        current_status = table_info_dict['status']
        dialog_title = f"Update Status: {table_info_dict['name']}" 
        
        possible_actions = []
        if current_status == "Kosong":
            possible_actions.append({"text": "BUAT PESANAN (BELUM BAYAR)", "new_status": "Belum Bayar"})
        elif current_status == "Belum Bayar":
            possible_actions.append({"text": "TANDAI SUDAH BAYAR", "new_status": "Sudah Bayar"})
            possible_actions.append({"text": "BATALKAN PESANAN (JADI KOSONG)", "new_status": "Kosong"})
        elif current_status == "Sudah Bayar":
            possible_actions.append({"text": "SIAPKAN MEJA BARU (JADI KOSONG)", "new_status": "Kosong"})
        
        dialog_buttons = [MDFlatButton(text="TUTUP", on_release=lambda x: self.status_update_dialog.dismiss())] #type: ignore 
        
        for action_details in possible_actions:
            btn = MDRaisedButton(
                text=action_details["text"],
                font_size="12sp"
            )
            btn.bind(on_release=lambda instance, new_s=action_details["new_status"]: self.confirm_status_update(new_s))
            dialog_buttons.append(btn)
        
        if not self.status_update_dialog:
            self.status_update_dialog = MDDialog(
                title=dialog_title,
                type="confirmation",
                buttons=dialog_buttons,
            )
        else:
            self.status_update_dialog.title = dialog_title
            self.status_update_dialog.buttons = dialog_buttons

        self.status_update_dialog.open()

    def confirm_status_update(self, new_status):
        if self.selected_table_for_status_update:
            table_name_to_update = self.selected_table_for_status_update['name']
            
            success = db_manager.update_table_status(table_name_to_update, new_status)
            
            if success:
                self.tables_data = db_manager.get_all_tables()
                self.display_tables()  
        
        if self.status_update_dialog:
            self.status_update_dialog.dismiss()
        self.selected_table_for_status_update = None

class Setting(MDScreen):
    pass

class MainApp(MDApp):
    root = None
    current_active = StringProperty("Add Menu")
    cart_items = {} 
    active_outlet = StringProperty("")
    last_outlet_name = StringProperty("")

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.theme_cls.primary_palette = 'Red'
        self.cart_items = {}

    def add_item_to_app_cart(self, menu_card_instance):
        item_name = menu_card_instance.name
        item_price = menu_card_instance.price
        item_image = menu_card_instance.actual_image_source

        if item_name in self.cart_items:
            self.cart_items[item_name]['quantity'] += 1
        else:
            self.cart_items[item_name] = {
                'name': item_name,
                'price': item_price,
                'quantity': 1,
                'image_path': item_image
            }
        if self.root and self.root.ids.screen_manager.current == 'Cart':
            cart_screen = self.root.ids.screen_manager.get_screen('Cart')
            cart_screen.update_cart_display()

    def update_cart_item_quantity(self, item_name, change_amount):
        if item_name in self.cart_items:
            self.cart_items[item_name]['quantity'] += change_amount
            if self.cart_items[item_name]['quantity'] <= 0:
                del self.cart_items[item_name]

        if self.root and self.root.ids.screen_manager.current == 'Cart':
            cart_screen = self.root.ids.screen_manager.get_screen('Cart')
            cart_screen.update_cart_display()

    def get_cart_total(self):
        total = 0.0
        for item_data in self.cart_items.values():
            total += item_data['price'] * item_data['quantity']
        return total

    def get_cart_items_list(self):
        return list(self.cart_items.values())

    def on_start(self):
        db_manager.init_db()
        outlet = db_manager.get_last_outlet()
        if outlet:
            self.active_outlet = outlet['name']
            self.last_outlet_name = outlet['name']
        else:
            self.active_outlet = "Default Outlet"
            self.last_outlet_name = "Default Outlet"

        if self.root:
            screen_manager = self.root.ids.get('screen_manager')
            if screen_manager:
                if not screen_manager.current:
                    screen_manager.current = 'Add Menu'
            else:
                print(f"[ERROR]: Screen manager tidak ditemukan.")

    def switch_screen(self, screen_name):
        if not self.root:
            return

        screen_manager = self.root.ids.get('screen_manager')
        if not screen_manager:
            return

        if screen_manager.current == screen_name and self.current_active == screen_name :
            return

        self.current_active = screen_name

        try:
            current_screen_name = screen_manager.current
            screen_names_list = screen_manager.screen_names

            if not current_screen_name or current_screen_name not in screen_names_list:
                screen_manager.transition = NoTransition()
                screen_manager.current = screen_name
                return

            if screen_manager.current == screen_name:
                return

            current_idx = screen_names_list.index(current_screen_name)
            target_idx = screen_names_list.index(screen_name)

            direction = ''
            if target_idx > current_idx:
                direction = 'left'
            elif target_idx < current_idx:
                direction = 'right'
            else:
                screen_manager.transition = NoTransition()
                screen_manager.current = screen_name
                return

            screen_manager.transition = SlideTransition(direction=direction, duration=0.2)
            screen_manager.current = screen_name

        except Exception as e:
            screen_manager.transition = NoTransition()
            screen_manager.current = screen_name

    def build(self):
        self.theme_cls.theme_style = "Light"
        return

    def events(self, window, key, *args):
        if key == 27: 
            if AddMenu.dialog and hasattr(AddMenu.dialog.content_cls, 'manager_open') and AddMenu.dialog.content_cls.manager_open:
                if hasattr(AddMenu.dialog.content_cls, 'file_manager') and AddMenu.dialog.content_cls.file_manager:
                    AddMenu.dialog.content_cls.file_manager.back()
                    return True
            elif AddMenu.dialog and AddMenu.dialog.open_button_press_time != 0:
                AddMenu.dialog.dismiss()
                return True
        return False

    def get_last_outlet_name(self):
        import db_manager
        outlet = db_manager.get_last_outlet()
        return outlet['name'] if outlet else ''

    def show_snackbar(self, message):
        Snackbar(text=message, duration=2).open()

    def show_dialog(self, title, text):
        dialog = MDDialog(title=title, text=text, size_hint=(0.8, None), height=dp(150))
        dialog.open()

    def export_sales_report_to_excel(self, outlet_name):
        from kivy.utils import platform
        import db_manager
        import os
        from datetime import datetime
        try:
            if not outlet_name.strip():
                self.show_snackbar('Nama outlet harus diisi!')
                return
            db_manager.save_outlet(outlet_name.strip())
            sales = db_manager.get_all_sales()
            if not sales:
                self.show_snackbar('Belum ada data penjualan!')
                return
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            if ws is not None:
                ws.title = 'Laporan Penjualan'
                ws.append(['Tanggal', 'Outlet', 'Item', 'Qty', 'Harga Satuan', 'Subtotal', 'Total'])
                for sale in sales:
                    for item in sale['items']:
                        ws.append([
                            sale['date'],
                            sale['outlet_name'],
                            item['name'],
                            item['quantity'],
                            item['price'],
                            item['price'] * item['quantity'],
                            sale['total']
                        ])
            else:
                self.show_dialog('Gagal', 'Gagal membuat worksheet Excel.')
                return
            filename = f"LaporanPenjualan_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            if platform == 'android':
                from android.storage import primary_external_storage_path # type: ignore
                download_dir = os.path.join(primary_external_storage_path(), 'Download')
            else:
                download_dir = os.path.expanduser('~/Download')
            if not os.path.exists(download_dir):
                os.makedirs(download_dir)
            file_path = os.path.join(download_dir, filename)
            try:
                wb.save(file_path)
                print(f"[DEBUG] File berhasil disimpan di: {file_path}")
                self.show_dialog('Sukses', f'Laporan berhasil disimpan:\n{file_path}')
            except Exception as e:
                print(f"[ERROR] Gagal menyimpan file: {e}")
                self.show_dialog('Error', f'Gagal menyimpan file: {e}')
            if platform == 'android':
                from jnius import autoclass # type: ignore
                PythonActivity = autoclass('org.kivy.android.PythonActivity')
                Intent = autoclass('android.content.Intent')
                Uri = autoclass('android.net.Uri')
                File = autoclass('java.io.File')
                context = PythonActivity.mActivity
                file = File(file_path)
                uri = Uri.fromFile(file)
                intent = Intent(Intent.ACTION_VIEW)
                intent.setDataAndType(uri, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                intent.addFlags(Intent.FLAG_GRANT_READ_URI_PERMISSION)
                context.startActivity(intent)
        except Exception as e:
            print(f"[ERROR] Gagal mengekspor laporan penjualan: {e}")
            
    def get_active_outlet(self):
        if not self.active_outlet:
            import db_manager
            outlet = db_manager.get_last_outlet()
            if outlet:
                self.active_outlet = outlet['name']
        return self.active_outlet

if __name__ == "__main__":
    MainApp().run()